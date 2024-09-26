from django.shortcuts import render, redirect
from django.conf import settings
from .forms import ConsultaForm, DocumentoForm
from .google_drive import search_files_in_drive, download_file, extract_text_from_file, get_drive_service
import logging
import os
from PIL import Image
import pytesseract
import fitz  # PyMuPDF
import re
import docx
import openpyxl
from datetime import datetime
import hashlib
from django.urls import reverse
from django.http import HttpResponse, JsonResponse
from .google_drive import get_drive_service, search_files_in_drive, vaciar_carpeta_drive, descargar_archivo
from django.contrib.auth.decorators import user_passes_test
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

# Configurar ruta al token
TOKEN_PATH = os.path.join(settings.BASE_DIR, 'token.json')

# Configura el logger
logger = logging.getLogger(__name__)

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def preprocess_image(image_path):
    try:
        image = Image.open(image_path)
        # Convertir la imagen a escala de grises
        image = image.convert('L')
        # Aplicar filtros adicionales si es necesario
        image = image.point(lambda x: 0 if x < 150 else 255, '1')  # Cambiar umbral según necesidad
        return image
    except Exception as e:
        logger.error(f"Error al preprocesar la imagen {image_path}: {e}")
        return None

def extract_text_from_image(image_path):
    try:
        # Asegúrate de que las rutas sean válidas
        image = Image.open(rf'{image_path}')
        # Configuraciones de Tesseract
        config = '--oem 3 --psm 6'
        text = pytesseract.image_to_string(image, config=config, lang='spa')
        return text
    except Exception as e:
        logger.error(f"Error al extraer texto de la imagen {image_path}: {e}")
        return ""

def extract_text_from_pdf(pdf_path):
    try:
        document = fitz.open(pdf_path)
        text = ""
        for page_num in range(document.page_count):
            page = document.load_page(page_num)
            text += page.get_text()
        return text
    except Exception as e:
        logger.error(f"Error al extraer texto del archivo PDF {pdf_path}: {e}")
        return ""

def extract_text_from_docx(docx_path):
    try:
        doc = docx.Document(docx_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        logger.error(f"Error al extraer texto del archivo DOCX {docx_path}: {e}")
        return ""

def extract_text_from_xlsx(xlsx_path):
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.extend([str(cell) for cell in row if cell is not None])
        return ' '.join(text)
    except Exception as e:
        logger.error(f"Error al extraer texto del archivo XLSX {xlsx_path}: {e}")
        return ""

def limpiar_texto(texto):
    """Elimina caracteres especiales y convierte el texto a minúsculas para una comparación más robusta."""
    return re.sub(r'[^A-Za-z0-9\s]', '', texto).strip().lower()

def buscar_termino(texto, consulta):
    """Realiza la búsqueda del término en el texto. Considera palabras similares o cercanas."""
    texto_limpio = limpiar_texto(texto)
    consulta_limpia = limpiar_texto(consulta)

    if consulta_limpia in texto_limpio:
        return True

    texto_palabras = texto_limpio.split()
    consulta_palabras = consulta_limpia.split()

    for palabra in consulta_palabras:
        if any(palabra in palabra_texto for palabra_texto in texto_palabras):
            return True

    return False

# Vista de consulta con búsqueda solo en modo local
def consulta_view(request):
    resultados = None
    cantidad_archivos = 0
    search_done = False

    if request.method == 'POST':
        form = ConsultaForm(request.POST)
        if form.is_valid():
            search_done = True
            consulta = form.cleaned_data['consulta']
            logger.info(f"Consulta: {consulta}")
            resultados = []

            archivos_dir = settings.ARCHIVOS_DIR
            for root, dirs, files in os.walk(archivos_dir):
                for file_name in files:
                    try:
                        archivo_path = os.path.join(root, file_name)
                        texto = ""

                        if file_name.lower().endswith('.pdf'):
                            texto = extract_text_from_pdf(archivo_path)
                        elif file_name.lower().endswith(('.png', '.jpg', '.jpeg')):
                            texto = extract_text_from_image(archivo_path)
                        elif file_name.lower().endswith('.docx'):
                            texto = extract_text_from_docx(archivo_path)
                        elif file_name.lower().endswith('.xlsx'):
                            texto = extract_text_from_xlsx(archivo_path)

                        if buscar_termino(texto, consulta):
                            # Verificar si el archivo está en la carpeta de documentos o descargados
                            ruta_archivo = obtener_ruta_archivo(file_name)
                            if ruta_archivo:
                                fecha_modificacion = datetime.fromtimestamp(os.path.getmtime(archivo_path)).strftime('%d/%m/%Y')
                                resultados.append({
                                    'nombre': file_name,
                                    'url': ruta_archivo,
                                    'fecha': fecha_modificacion
                                })
                    except Exception as e:
                        logger.error(f"Error procesando archivo {file_name}: {e}")

            cantidad_archivos = len(resultados)

    else:
        form = ConsultaForm()

    context = {
        'form': form,
        'resultados': resultados if resultados else None,
        'cantidad_archivos': cantidad_archivos,
        'search_done': search_done
    }
    return render(request, 'documentos/consulta.html', context)

def subir_comprobante_view(request):
    if request.method == 'POST':
        form = DocumentoForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('consulta')
    else:
        form = DocumentoForm()
    return render(request, 'documentos/subir_comprobante.html', {'form': form})

def archivo_existe(ruta_descarga, nombre_archivo):
    archivo_path = os.path.join(ruta_descarga, nombre_archivo)
    return os.path.exists(archivo_path)

# Verifica si el archivo ya existe en la ruta de descarga o en la carpeta principal de documentos.
def obtener_ruta_archivo(file_name):
    # Ruta en /media/documentos
    ruta_documentos = os.path.join(settings.MEDIA_ROOT, 'documentos', file_name)
    
    # Ruta en /media/documentos/descargados
    ruta_descargados = os.path.join(settings.MEDIA_ROOT, 'documentos', 'descargados', file_name)
    
    # Verificar si el archivo existe en alguna de las dos rutas
    if os.path.exists(ruta_documentos):
        return os.path.join(settings.MEDIA_URL, 'documentos', file_name)
    elif os.path.exists(ruta_descargados):
        return os.path.join(settings.MEDIA_URL, 'documentos', 'descargados', file_name)
    else:
        return None
    

# Nueva función: Descargar archivos desde Google Drive
def descargar_archivos_nube(request):
    try:
        drive_files = search_files_in_drive(settings.DRIVE_FOLDER_ID)
        service = get_drive_service()

        if not drive_files:
            logger.info("No se encontraron archivos en la carpeta de Google Drive.")
            return render(request, 'documentos/sin_archivos.html')

        ruta_descarga = os.path.join(settings.MEDIA_ROOT, 'documentos', 'descargados')
        if not os.path.exists(ruta_descarga):
            os.makedirs(ruta_descarga)

        archivos_descargados = []
        archivos_omitidos = []

        for file in drive_files:
            try:
                file_id = file['id']
                file_name = file['name']
                file_data = download_file(service, file_id)

                if file_data:
                    created_time = file.get('createdTime', '')

                    if created_time:
                        try:
                            fecha_formateada = datetime.strptime(created_time, "%Y-%m-%dT%H:%M:%S.%fZ").strftime("%Y%m%d_%H%M%S")
                        except ValueError:
                            logger.error(f"Error al formatear la fecha {created_time}. Asignando 'Fecha desconocida'.")
                            fecha_formateada = "Fecha_desconocida"
                    else:
                        logger.warning(f"El archivo {file_name} no tiene 'createdTime'. Asignando 'Fecha desconocida'.")
                        fecha_formateada = "Fecha_desconocida"

                    extension = os.path.splitext(file_name)[1]
                    id_corto = file_id[:8]
                    nuevo_nombre = f"bethania-{fecha_formateada}-{id_corto}{extension}"

                    if archivo_existe(ruta_descarga, nuevo_nombre):
                        logger.info(f"El archivo {nuevo_nombre} ya existe. Omitiendo descarga.")
                        archivos_omitidos.append(nuevo_nombre)
                    else:
                        archivo_path = os.path.join(ruta_descarga, nuevo_nombre)
                        with open(archivo_path, 'wb') as archivo_local:
                            archivo_local.write(file_data)
                        logger.info(f"Archivo {file_name} descargado exitosamente como {nuevo_nombre}.")
                        archivos_descargados.append(nuevo_nombre)

                else:
                    logger.error(f"No se pudo obtener el contenido del archivo {file_name}.")
            except Exception as e:
                logger.error(f"Error descargando archivo {file_name}: {e}")
                return render(request, 'documentos/error_descarga.html', {'mensaje_error': f"Error descargando archivo {file_name}: {e}"})

        return render(request, 'documentos/resumen_descarga.html', {
            'archivos_descargados': archivos_descargados,
            'archivos_omitidos': archivos_omitidos
        })
    except Exception as e:
        logger.error(f"Error en la descarga de archivos: {e}")
        return render(request, 'documentos/error_descarga.html', {'mensaje_error': f"Error en la descarga de archivos: {e}"})

def es_superadministrador(user):
    """Verifica si el usuario es un superadministrador."""
    return user.is_superuser

# Nueva función: Vaciar la carpeta de Google Drive (solo para superadministradores)
@user_passes_test(es_superadministrador, login_url='/forbidden/')
def vaciar_carpeta_drive(request):
    try:
        service = get_drive_service()
        folder_id = settings.DRIVE_FOLDER_ID
        results = service.files().list(q=f"'{folder_id}' in parents").execute()
        files = results.get('files', [])

        if not files:
            return render(request, 'documentos/carpeta_vacia.html')

        for file in files:
            try:
                service.files().delete(fileId=file['id']).execute()
            except Exception as error:
                logger.error(f"Error al eliminar el archivo {file['name']}: {error}")

        return render(request, 'documentos/carpeta_vaciada_exito.html')
    except Exception as e:
        logger.error(f"Error vaciando la carpeta: {e}")
        return HttpResponse(f"Error vaciando la carpeta: {e}")

def forbidden_view(request):
    """Muestra un mensaje de error cuando el usuario no tiene permisos suficientes."""
    return render(request, 'forbidden.html')
