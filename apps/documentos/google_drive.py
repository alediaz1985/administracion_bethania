import logging
from io import BytesIO
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from PyPDF2 import PdfReader
from PIL import Image
import pytesseract
import docx
import openpyxl
from django.conf import settings

# Configura la ruta a tu archivo de credenciales
CREDENTIALS_FILE = settings.GOOGLE_CREDENTIALS

SCOPES = ['https://www.googleapis.com/auth/drive']

# Configura el logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def get_drive_service():
    try:
        creds = service_account.Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
        service = build('drive', 'v3', credentials=creds)
        logger.info("Autenticación con Google Drive exitosa")
        return service
    except Exception as e:
        logger.error(f"Error en la autenticación con Google Drive: {e}")
        return None

def search_files_in_drive(folder_id, mime_types=None):
    service = get_drive_service()
    if not service:
        return []

    try:
        query = f"'{folder_id}' in parents"
        if mime_types:
            mime_query = " or ".join([f"mimeType='{mime}'" for mime in mime_types])
            query += f" and ({mime_query})"
        logger.info(f"Query utilizada para la búsqueda: {query}")
        print(f"Query utilizada para la búsqueda: {query}")

        results = service.files().list(
            q=query,
            pageSize=1000,
            fields="nextPageToken, files(id, name, mimeType, createdTime, webViewLink)"
        ).execute()
        items = results.get('files', [])
        logger.info(f"Archivos encontrados: {len(items)}")
        print(f"Archivos encontrados: {len(items)}")
        for item in items:
            logger.info(f"Archivo encontrado: {item['name']} (ID: {item['id']}, MIME type: {item['mimeType']})")
            print(f"Archivo encontrado: {item['name']} (ID: {item['id']}, MIME type: {item['mimeType']})")
        return items
    except Exception as e:
        logger.error(f"Error buscando archivos en Google Drive: {e}")
        return []

def download_file(service, file_id):
    try:
        request = service.files().get_media(fileId=file_id)
        file_data = BytesIO()  # Creamos un BytesIO para almacenar los datos descargados
        downloader = MediaIoBaseDownload(file_data, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
        
        file_data.seek(0)  # Nos aseguramos de posicionarnos al principio del archivo
        
        return file_data.read()  # Aquí devolvemos los datos en formato de bytes
    except Exception as e:
        logger.error(f"Error descargando archivo: {e}")
        return None

def extract_text_from_pdf(file_data):
    try:
        pdf = PdfReader(BytesIO(file_data))
        text = ""
        for page in pdf.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        logger.error(f"Error al extraer texto del PDF: {e}")
        return ""

def extract_text_from_image(file_data):
    try:
        image = Image.open(BytesIO(file_data))
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        logger.error(f"Error al extraer texto de la imagen: {e}")
        return ""

def extract_text_from_docx(file_data):
    try:
        doc = docx.Document(BytesIO(file_data))
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        logger.error(f"Error al extraer texto del DOCX: {e}")
        return ""

def extract_text_from_xlsx(file_data):
    try:
        wb = openpyxl.load_workbook(BytesIO(file_data), data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.extend([str(cell) for cell in row if cell is not None])
        return ' '.join(text)
    except Exception as e:
        logger.error(f"Error al extraer texto del XLSX: {e}")
        return ""

def extract_text_from_file(file_data, mime_type):
    if mime_type == 'application/pdf':
        return extract_text_from_pdf(file_data)
    elif mime_type in ['image/png', 'image/jpeg']:
        return extract_text_from_image(file_data)
    elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        return extract_text_from_docx(file_data)
    elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        return extract_text_from_xlsx(file_data)
    else:
        return ""

def vaciar_carpeta_drive(folder_id):
    try:
        service = get_drive_service()
        if not service:
            return "Error al autenticar con Google Drive"

        # Listar todos los archivos en la carpeta
        archivos = service.files().list(q=f"'{folder_id}' in parents").execute().get('files', [])

        if not archivos:
            logger.info("No se encontraron archivos en la carpeta de Google Drive.")
            return "No se encontraron archivos en la carpeta."

        # Eliminar todos los archivos encontrados
        for archivo in archivos:
            try:
                file_id = archivo['id']
                service.files().delete(fileId=file_id).execute()
                logger.info(f"Archivo {archivo['name']} eliminado exitosamente.")
            except Exception as e:
                logger.error(f"Error eliminando archivo {archivo['name']}: {e}")

        return "Todos los archivos fueron eliminados exitosamente."
    except Exception as e:
        logger.error(f"Error vaciando la carpeta de Google Drive: {e}")
        return f"Error vaciando la carpeta: {e}"